#!/usr/bin/env python3
"""
Compliance Q&A: reads PDFs from a folder, answers questions from a spreadsheet
using the Anthropic API, and writes answers back to the spreadsheet.

Usage:
    uv run main.py <questions.xlsx> [pdf_dir]

Spreadsheet format:
    Row 1:   Header row  (column A = questions header, column B = answers header)
    Rows 2+: One question per row in column A

PDF content is cached across all questions to minimise API cost. Rows that
already have an answer in column B are skipped, so an interrupted run can be
resumed by simply re-running (skipped/blank rows are asked again).
"""

import sys
from pathlib import Path

import anthropic
import pypdf
import openpyxl
from dotenv import load_dotenv

load_dotenv()

MODEL = "claude-sonnet-5-5"
MAX_ANSWER_TOKENS = 512
# If an answer hits the token cap it's truncated mid-sentence; retry once with
# this larger cap before flagging the cell.
RETRY_ANSWER_TOKENS = 2048
DEFAULT_PDF_DIR = "pdfs"

SYSTEM_INSTRUCTION = """You are a compliance assistant answering questions from policy documents.

Rules:
- Answer directly and concisely. No preamble, no "Based on the policy documents..." or similar phrases.
- If the answer is clearly present in the documents, state it plainly.
- If the question asks for a document to be provided or attached, respond with exactly: SKIP
- If the answer cannot be found in the documents, respond with exactly: SKIP
- Do not explain why you are skipping. Do not say "I cannot find..." or "The documents do not...". Just: SKIP"""


def extract_pdf_text(pdf_path: Path) -> str:
    reader = pypdf.PdfReader(str(pdf_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return f"=== {pdf_path.name} ===\n\n" + "\n\n".join(pages)


def load_documents(pdf_dir: Path) -> str:
    pdf_files = sorted(pdf_dir.glob("*.pdf"))
    if not pdf_files:
        print(f"Error: no PDF files found in '{pdf_dir}'.")
        print("Copy your policy PDFs into that folder and re-run.")
        sys.exit(1)

    print(f"Loading {len(pdf_files)} PDF(s):")
    docs = []
    for pdf_path in pdf_files:
        print(f"  {pdf_path.name}")
        try:
            docs.append(extract_pdf_text(pdf_path))
        except Exception as exc:
            print(f"  Warning: could not read {pdf_path.name}: {exc}")

    if not docs:
        print("Error: all PDFs failed to load.")
        sys.exit(1)

    separator = "\n\n" + "=" * 60 + "\n\n"
    return separator.join(docs)


def answer_questions(spreadsheet_path: Path, pdf_dir: Path) -> None:
    client = anthropic.Anthropic()

    documents = load_documents(pdf_dir)

    wb = openpyxl.load_workbook(spreadsheet_path)
    ws = wb.active

    if ws.max_row < 2:
        print("Error: spreadsheet needs at least 2 rows (header + questions).")
        sys.exit(1)

    # Write answer column header
    ws["B1"] = "Answer"

    # Collect non-empty questions from column A (skip header row). Rows whose
    # answer cell is already filled were answered on a previous run.
    questions: list[tuple[int, str]] = []
    answered = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=False):
        cell = row[0]
        if cell.value and str(cell.value).strip():
            existing = ws.cell(row=cell.row, column=2).value
            if existing is not None and str(existing).strip():
                answered += 1
            else:
                questions.append((cell.row, str(cell.value).strip()))

    if answered:
        print(f"Skipping {answered} already-answered row(s).")
    if not questions:
        print("Nothing to do: no unanswered questions in column A.")
        return

    print(f"\nAnswering {len(questions)} question(s) using {MODEL}...\n")

    for i, (row_num, question) in enumerate(questions, 1):
        label = question[:80] + ("..." if len(question) > 80 else "")
        print(f"[{i}/{len(questions)}] {label}")

        def ask(max_tokens: int) -> "anthropic.types.Message":
            return client.messages.create(
                model=MODEL,
                max_tokens=max_tokens,
                system=[
                    # Cache the system instruction so it stays at the front of
                    # the prefix and is reused across all questions.
                    {
                        "type": "text",
                        "text": SYSTEM_INSTRUCTION,
                        "cache_control": {"type": "ephemeral"},
                    },
                    # Cache the full document corpus — identical for every question.
                    {
                        "type": "text",
                        "text": f"Policy Documents:\n\n{documents}",
                        "cache_control": {"type": "ephemeral"},
                    },
                ],
                messages=[{"role": "user", "content": question}],
            )

        response = ask(MAX_ANSWER_TOKENS)
        if response.stop_reason == "max_tokens":
            print("  Answer hit the token cap — retrying with a larger one")
            response = ask(RETRY_ANSWER_TOKENS)

        answer = next(
            (block.text for block in response.content if block.type == "text"), ""
        ).strip()
        if response.stop_reason == "max_tokens":
            answer += " [TRUNCATED]"

        # Leave blank for skipped questions
        ws.cell(row=row_num, column=2).value = (
            "" if answer.upper().startswith("SKIP") else answer
        )

        # Save after every answer so progress survives an interruption.
        wb.save(spreadsheet_path)

        # Log cache activity on the first call to confirm caching is working.
        if i == 1:
            usage = response.usage
            written = getattr(usage, "cache_creation_input_tokens", 0)
            read = getattr(usage, "cache_read_input_tokens", 0)
            if written:
                print(f"  Cache written: {written:,} tokens")
            if read:
                print(f"  Cache hit:     {read:,} tokens")

    print(f"\nDone. Answers written to {spreadsheet_path}")


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    spreadsheet_path = Path(sys.argv[1])
    pdf_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(DEFAULT_PDF_DIR)

    if not spreadsheet_path.exists():
        print(f"Error: '{spreadsheet_path}' not found.")
        sys.exit(1)

    if not pdf_dir.is_dir():
        print(f"Error: PDF directory '{pdf_dir}' not found.")
        sys.exit(1)

    answer_questions(spreadsheet_path, pdf_dir)


if __name__ == "__main__":
    main()
